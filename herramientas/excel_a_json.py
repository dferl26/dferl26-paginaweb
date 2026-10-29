#!/usr/bin/env python3
"""
Dayvilab — Convertidor de catálogo Excel/CSV a productos.json

USO
  Excel (recomendado, usa la plantilla plantilla-productos.xlsx):
      python3 excel_a_json.py catalogo.xlsx
  CSV (exportado desde Excel; acepta ";" o ","):
      python3 excel_a_json.py productos.csv --especificaciones especificaciones.csv --presentaciones presentaciones.csv
  Opciones:
      -o, --salida RUTA   Archivo de salida (por defecto: ../productos.json)
      --legible           JSON con sangría (más pesado; por defecto se genera compacto)
      --forzar            Genera el archivo aunque haya filas con errores (esas filas se omiten)

REQUISITOS
  Python 3.8+  ·  Para .xlsx:  pip install openpyxl

COLUMNAS DE LA HOJA "productos"
  id* · codigo · nombre* · categoria* · subcategoria · marca · modelo · resumen · descripcion
  imagen · imagenes_adicionales (URLs separadas por |) · etiquetas (separadas por ,)
  destacado (si/no) · ficha_tecnica_url · hoja_seguridad_url
  esp: <Nombre>   → (opcional) columnas que empiezan con "esp:" se agregan a la tabla de especificaciones
                    (ej. "esp: Capacidad"). Útil para catálogos homogéneos.

HOJA "especificaciones" (recomendada; o archivo --especificaciones en CSV)
  producto_id* · atributo* · valor*      → una fila por especificación, en el orden en que se mostrará

HOJA "presentaciones" (opcional; o archivo --presentaciones en CSV)
  producto_id* · codigo · descripcion* · unidad

categoria acepta: equipos · materiales (cristalería/plásticos) · reactivos · medico (material médico)
"""
import argparse
import csv
import datetime as dt
import io
import json
import re
import sys
import unicodedata
from pathlib import Path

CATEGORIAS = {
    'equipos': 'equipos', 'equipo': 'equipos', 'equipos de laboratorio': 'equipos',
    'materiales': 'materiales', 'material': 'materiales', 'cristaleria': 'materiales', 'plasticos': 'materiales',
    'cristaleria y materiales': 'materiales', 'cristaleria, plasticos y materiales': 'materiales',
    'cristaleria/plasticos': 'materiales',
    'reactivos': 'reactivos', 'reactivo': 'reactivos', 'consumibles': 'reactivos',
    'reactivos y consumibles': 'reactivos', 'reactivos y consumibles analiticos': 'reactivos',
    'medico': 'medico', 'material medico': 'medico', 'insumos medicos': 'medico',
}
SI = {'si', 'sí', 's', 'x', '1', 'true', 'verdadero', 'yes'}


def sin_tildes(s):
    return ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn')


def clave(s):
    """Normaliza un encabezado: minúsculas, sin tildes, espacios simples."""
    return re.sub(r'\s+', ' ', sin_tildes(str(s or '')).strip().lower())


def texto(v):
    """Convierte el valor de una celda a texto limpio (250.0 → '250', fechas → ISO)."""
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (dt.datetime, dt.date)):
        return v.date().isoformat() if isinstance(v, dt.datetime) else v.isoformat()
    return str(v).strip()


def id_valido(s):
    """IDs aptos para URL: letras, números, guion, guion bajo y punto."""
    s2 = re.sub(r'\s+', '-', sin_tildes(s.strip()))
    s2 = re.sub(r'[^A-Za-z0-9._-]', '', s2)
    return s2


# ---------- Lectura ----------
def leer_xlsx(ruta):
    try:
        import openpyxl
    except ImportError:
        sys.exit('Falta openpyxl. Instálelo con:  pip install openpyxl')
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    hojas = {clave(n): n for n in wb.sheetnames}
    nombre_prod = hojas.get('productos', wb.sheetnames[0])
    filas = lambda nombre: [list(r) for r in wb[nombre].iter_rows(values_only=True)]
    productos = filas(nombre_prod)
    presentaciones = filas(hojas['presentaciones']) if 'presentaciones' in hojas else []
    especificaciones = filas(hojas['especificaciones']) if 'especificaciones' in hojas else []
    return productos, especificaciones, presentaciones


def leer_csv(ruta):
    raw = Path(ruta).read_bytes()
    for enc in ('utf-8-sig', 'cp1252'):  # Excel en Windows suele exportar en cp1252
        try:
            data = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        delim = csv.Sniffer().sniff(data[:5000], delimiters=';,\t').delimiter
    except csv.Error:
        delim = ';' if data.count(';') > data.count(',') else ','
    return [row for row in csv.reader(io.StringIO(data), delimiter=delim)]


def a_dicts(filas):
    filas = [f for f in filas if f and any(texto(c) for c in f)]
    if not filas:
        return [], []
    enc = [texto(h) for h in filas[0]]
    salida = []
    for n, f in enumerate(filas[1:], start=2):
        d = {enc[i]: texto(f[i]) if i < len(f) else '' for i in range(len(enc)) if enc[i]}
        d['_fila'] = n
        salida.append(d)
    return enc, salida


# ---------- Conversión ----------
def agrupar(filas, requeridas, nombre_hoja, avisos):
    """Lee una hoja relacional (producto_id + columnas) y agrupa sus filas por producto."""
    if not filas:
        return {}, {}
    enc, datos = a_dicts(filas)
    m = {clave(h): h for h in enc if h}
    faltan = [r for r in requeridas if r not in m]
    if faltan:
        avisos.append(f'{nombre_hoja}: faltan las columnas {faltan}; se omitió la hoja.')
        return {}, {}
    grupos = {}
    for d in datos:
        pid = id_valido(d.get(m['producto_id'], ''))
        if pid and all(d.get(m[r], '') for r in requeridas):
            grupos.setdefault(pid, []).append(d)
    return grupos, m


def convertir(prod_rows, esp_rows, pres_rows):
    errores, avisos = [], []
    encabezados, filas = a_dicts(prod_rows)
    mapa = {clave(h): h for h in encabezados if h}
    for req in ('id', 'nombre', 'categoria'):
        if req not in mapa:
            sys.exit(f'La hoja de productos no tiene la columna obligatoria "{req}". Encabezados: {encabezados}')
    col = lambda d, k: d.get(mapa.get(k, ''), '')
    esp_cols = [h for h in encabezados if h and clave(h).startswith('esp:')]

    # Hojas relacionales agrupadas por producto
    grupos_pres, pm = agrupar(pres_rows, ['producto_id', 'descripcion'], 'Presentaciones', avisos)
    pres_por_id = {}
    for pid, ds in grupos_pres.items():
        for d in ds:
            item = {'descripcion': d[pm['descripcion']]}
            for k in ('codigo', 'unidad'):
                v = d.get(pm.get(k, ''), '')
                if v:
                    item[k] = v
            pres_por_id.setdefault(pid, []).append(item)
    grupos_esp, em = agrupar(esp_rows, ['producto_id', 'atributo', 'valor'], 'Especificaciones', avisos)
    esp_por_id = {pid: [(d[em['atributo']], d[em['valor']]) for d in ds] for pid, ds in grupos_esp.items()}

    productos, vistos = [], set()
    for d in filas:
        fila = d['_fila']
        id_orig = col(d, 'id')
        nombre = col(d, 'nombre')
        if not id_orig and not nombre:
            continue
        pid = id_valido(id_orig)
        if not pid:
            errores.append(f'Fila {fila}: falta el id.')
            continue
        if pid != id_orig:
            avisos.append(f'Fila {fila}: id "{id_orig}" se normalizó a "{pid}" para usarlo en la URL.')
        if pid in vistos:
            errores.append(f'Fila {fila}: id duplicado "{pid}".')
            continue
        if not nombre:
            errores.append(f'Fila {fila}: falta el nombre del producto "{pid}".')
            continue
        cat = CATEGORIAS.get(clave(col(d, 'categoria')))
        if not cat:
            errores.append(f'Fila {fila}: categoría "{col(d, "categoria")}" no válida (use equipos, materiales, reactivos o medico).')
            continue
        vistos.add(pid)

        p = {'id': pid, 'nombre': nombre, 'categoria': cat}
        for k in ('codigo', 'subcategoria', 'marca', 'modelo', 'resumen', 'descripcion', 'imagen'):
            v = col(d, k)
            if v:
                p[k] = v
        extra = [u.strip() for u in col(d, 'imagenes_adicionales').split('|') if u.strip()]
        if extra:
            p['imagenes'] = extra
        tags = [t.strip() for t in col(d, 'etiquetas').split(',') if t.strip()]
        if tags:
            p['etiquetas'] = tags
        if clave(col(d, 'destacado')) in SI:
            p['destacado'] = True
        esp = {}
        for atributo, valor in esp_por_id.pop(pid, []):
            esp[atributo] = valor
        for h in esp_cols:
            v = d.get(h, '')
            if v:
                esp[h.split(':', 1)[1].strip()] = v
        if esp:
            p['especificaciones'] = esp
        if pid in pres_por_id:
            p['presentaciones'] = pres_por_id.pop(pid)
        docs = []
        if col(d, 'ficha_tecnica_url'):
            docs.append({'nombre': 'Ficha técnica', 'url': col(d, 'ficha_tecnica_url')})
        if col(d, 'hoja_seguridad_url'):
            docs.append({'nombre': 'Hoja de seguridad (SDS)', 'url': col(d, 'hoja_seguridad_url')})
        if docs:
            p['documentos'] = docs
        if not p.get('imagen'):
            avisos.append(f'Fila {fila}: "{pid}" no tiene imagen; se mostrará el ícono de su categoría.')
        productos.append(p)

    for pid in pres_por_id:
        avisos.append(f'Presentaciones: el producto_id "{pid}" no existe en la hoja de productos.')
    for pid in esp_por_id:
        avisos.append(f'Especificaciones: el producto_id "{pid}" no existe en la hoja de productos.')
    return productos, errores, avisos


def main():
    ap = argparse.ArgumentParser(description='Convierte el catálogo de Dayvilab (Excel/CSV) a productos.json')
    ap.add_argument('entrada', help='Archivo .xlsx o .csv con los productos')
    ap.add_argument('--especificaciones', help='CSV de especificaciones (solo si la entrada es CSV)')
    ap.add_argument('--presentaciones', help='CSV de presentaciones (solo si la entrada es CSV)')
    ap.add_argument('-o', '--salida', default=str(Path(__file__).resolve().parent.parent / 'productos.json'))
    ap.add_argument('--legible', action='store_true', help='JSON con sangría')
    ap.add_argument('--forzar', action='store_true', help='Generar aunque haya errores')
    a = ap.parse_args()

    ruta = Path(a.entrada)
    if not ruta.exists():
        sys.exit(f'No se encontró el archivo: {ruta}')
    if ruta.suffix.lower() in ('.xlsx', '.xlsm'):
        prod_rows, esp_rows, pres_rows = leer_xlsx(ruta)
    elif ruta.suffix.lower() in ('.csv', '.txt'):
        prod_rows = leer_csv(ruta)
        esp_rows = leer_csv(a.especificaciones) if a.especificaciones else []
        pres_rows = leer_csv(a.presentaciones) if a.presentaciones else []
    else:
        sys.exit('Formato no soportado. Use .xlsx o .csv')

    productos, errores, avisos = convertir(prod_rows, esp_rows, pres_rows)
    for m in avisos:
        print('  aviso  ·', m)
    for m in errores:
        print('  ERROR  ·', m)
    if errores and not a.forzar:
        sys.exit(f'\n{len(errores)} error(es). Corrija el archivo o use --forzar para omitir esas filas.')

    salida = {
        'version': 1,
        'actualizado': dt.date.today().isoformat(),
        'total': len(productos),
        'productos': productos,
    }
    txt = json.dumps(salida, ensure_ascii=False, indent=2 if a.legible else None,
                     separators=None if a.legible else (',', ':'))
    Path(a.salida).write_text(txt + '\n', encoding='utf-8')
    por_cat = {}
    for p in productos:
        por_cat[p['categoria']] = por_cat.get(p['categoria'], 0) + 1
    print(f'\n✓ {len(productos)} productos → {a.salida} ({len(txt.encode()) / 1024:.0f} KB)')
    print('  ' + ' · '.join(f'{k}: {v}' for k, v in sorted(por_cat.items())))


if __name__ == '__main__':
    main()
